from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = PROJECT_ROOT / "docs"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), LIGHT_GRAY)


def configure_document(doc, subtitle):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.03

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(26)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style_properties = title_style._element.get_or_add_pPr()
    title_style_border = title_style_properties.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_properties.remove(title_style_border)

    for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        if name in styles:
            style = styles[name]
            style.font.name = "Aptos Display"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(doc.core_properties.title)
    title_properties = title._p.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = intro.add_run(subtitle)
    run.italic = True
    run.font.size = Pt(11)
    intro.paragraph_format.space_after = Pt(16)


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_module_table(doc):
    rows = [
        ("Interface", "index.html and style.css", "Camera view, status, calibration control, and session metrics"),
        ("Tracking", "script.js and MediaPipe Pose", "Estimates landmarks and draws the live skeleton"),
        ("Validation", "isRealHumanUpperBody", "Rejects frames without plausible face and shoulder geometry"),
        ("Classification", "predictPosture", "Compares normalized geometry with the active baseline"),
        ("Reference data", "data/hf_posture_dataset.csv", "Supports exploration but does not train the live classifier"),
        ("Documentation", "docs and tools/generate_docs.py", "Keeps project reports reproducible"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = (Inches(1.2), Inches(2.05), Inches(3.65))
    headers = ("Module", "Files or function", "Responsibility")
    for index, (cell, text, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for cell, text, width in zip(cells, row_data, widths):
            cell.width = width
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
    set_table_borders(table)


def create_conference_report():
    doc = Document()
    doc.core_properties.title = "Posture Monitor Conference Review Report"
    doc.core_properties.subject = "Implementation and architecture review"
    configure_document(
        doc,
        "Conference Paper Format: Architecture, Implementation, and Evaluation",
    )

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "This paper presents Posture Monitor, a static web application that performs real-time pose estimation "
        "and posture classification entirely in the browser. By leveraging MediaPipe Pose and scale-independent "
        "geometric feature extraction, the system delivers immediate visual feedback without relying on an "
        "application server. This privacy-first approach ensures that sensitive camera frames remain on the local "
        "device. The paper details the system architecture, geometric classification algorithms, validation "
        "methodology, and current limitations."
    )

    doc.add_heading("Keywords", level=1)
    doc.add_paragraph(
        "Posture Detection, MediaPipe Pose, WebRTC, Edge Computing, Privacy-First, Human-Computer Interaction"
    )

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "Maintaining proper posture is a critical factor in mitigating musculoskeletal disorders associated with "
        "prolonged desktop computer use. Traditional posture monitoring systems often require dedicated hardware "
        "or stream video to remote servers for processing, raising privacy concerns. This project addresses these "
        "issues by executing all pose estimation and classification algorithms within the client's browser."
    )

    doc.add_heading("2. System Architecture", level=1)
    doc.add_paragraph(
        "The architecture is designed to minimize latency and ensure data privacy. The data flow is strictly local: "
        "Camera Frame -> MediaPipe Pose -> Upper Body Validation -> Feature Calculation -> Threshold Classification -> UI."
    )
    doc.add_paragraph(
        "The browser owns the complete runtime path. It requests a single camera stream via WebRTC, displays a "
        "preview while MediaPipe initializes, and processes frames through a sequential requestAnimationFrame loop. "
        "This prevents duplicate camera streams and overlapping pose request bottlenecks."
    )

    doc.add_heading("3. Implementation Methodology", level=1)
    add_module_table(doc)

    doc.add_heading("3.1 Geometric Classification", level=2)
    doc.add_paragraph(
        "The application derives two scale-independent ratios to ensure robustness across varying camera distances. "
        "Feature f1 is the absolute difference between the left and right face-to-shoulder distances, scaled by ten "
        "and normalized by shoulder width. Feature f2 is the sum of those face-to-shoulder distances divided by "
        "shoulder width."
    )
    doc.add_paragraph(
        "A user-initiated calibration step stores the active f1 and f2 values as the baseline. The classifier triggers "
        "a 'slouch' state when f2 falls below the baseline minus a threshold (0.12), and a 'lean' state when f1 "
        "deviates from the baseline by more than 0.55."
    )

    validation_heading = doc.add_heading("3.2 Human Validation", level=2)
    validation_heading.paragraph_format.page_break_before = True
    add_bullets(
        doc,
        [
            "Verification of high confidence scores for nose, eyes, and shoulders.",
            "Requirement of at least one visible ear for geometric incenter calculation.",
            "Positional check ensuring shoulders appear vertically below the nose.",
            "Rejection of implausibly narrow shoulder spans (minimum width threshold).",
            "Verification of the head-to-shoulder distance ratio to filter false positives.",
        ],
    )

    doc.add_heading("4. Evaluation and Demonstration", level=1)
    doc.add_paragraph(
        "The system has been evaluated through manual user testing. The demonstration protocol involves serving the "
        "application over HTTPS, allowing camera access, and navigating through upright, slouched, and lateral lean "
        "states after establishing a baseline. Transition counters successfully increment only upon state changes "
        "rather than per-frame."
    )

    doc.add_heading("5. Limitations and Future Work", level=1)
    add_bullets(
        doc,
        [
            "Classification accuracy is dependent on environmental lighting, camera angle, and clothing contrast.",
            "Current heuristic thresholds (0.12 for slouch, 0.55 for lean) require broader clinical validation.",
            "Reliance on third-party CDNs for MediaPipe models introduces external dependencies.",
            "Future iterations should include time-based smoothing to filter brief, natural movements.",
        ],
    )

    doc.add_heading("6. Conclusion", level=1)
    doc.add_paragraph(
        "Posture Monitor successfully demonstrates that robust, real-time posture tracking can be achieved in a "
        "static web environment without compromising user privacy. The integration of scale-normalized geometric "
        "features and user-driven calibration provides a flexible and explainable classification model."
    )

    doc.save(DOCS_DIR / "CONFERENCE_REPORT_v2.docx")


def create_project_thesis():
    doc = Document()
    doc.core_properties.title = "Posture Monitor Engineering Report"
    doc.core_properties.subject = "Technical project report"
    configure_document(
        doc,
        "Design implementation, detailed code breakdown, validation, and limitations of a browser posture monitor",
    )

    doc.add_heading("Project Overview", level=1)
    doc.add_paragraph(
        "Posture Monitor demonstrates how a browser can combine webcam access, pose estimation, "
        "geometric feature extraction, and immediate visual feedback without an application "
        "server. The result is a small deployable system with a transparent classification path. "
        "Its main engineering contribution is the integration of MediaPipe landmarks with "
        "calibrated normalized measurements and a guarded real-time processing loop."
    )

    doc.add_heading("Problem Definition", level=1)
    doc.add_paragraph(
        "A useful desktop posture aid must respond quickly, adapt to camera placement, and avoid "
        "sending sensitive video to a custom backend. A single fixed neck angle is fragile because "
        "body proportions, seating position, lens height, and screen tilt vary. The project "
        "therefore uses ratios normalized by shoulder width and lets the user establish an upright "
        "baseline."
    )

    doc.add_heading("Architecture", level=1)
    doc.add_paragraph(
        "The application is delivered as HTML CSS and JavaScript. After the user grants camera "
        "permission, one MediaStream supplies the hidden video element. MediaPipe Pose processes "
        "frames sequentially. The result callback validates landmark geometry, calculates the two "
        "posture features, assigns a state, and renders the camera frame and skeleton on a canvas."
    )
    
    doc.add_heading("Detailed Code Explanation", level=1)
    doc.add_paragraph(
        "The project relies on three core files: index.html for structure, style.css for presentation, "
        "and script.js for logic. Below is a detailed breakdown of each code component."
    )

    doc.add_heading("index.html: Structure and Layout", level=2)
    doc.add_paragraph(
        "The index.html file provides the skeleton of the application. It includes:"
    )
    add_bullets(
        doc,
        [
            "External Libraries: Imports MediaPipe Pose and Drawing Utils from public CDNs using <script> tags in the <head>.",
            "Video & Canvas Elements: A hidden <video id='input_video'> element captures the raw webcam feed. A <canvas id='output_canvas'> element acts as the primary display, overlaying the processed skeleton over the video.",
            "Status Overlay & Calibration: A div (id='status-overlay') provides real-time text feedback (e.g., 'Slouching detected'). The 'calibrate-btn' button allows users to set their baseline posture.",
            "Stats Dashboard: A dedicated div container ('stats-dashboard') displays real-time metrics, including Session Time, Slouch Count, Lean Count, and Calibration status.",
        ],
    )

    doc.add_heading("style.css: Styling and Responsive Design", level=2)
    doc.add_paragraph(
        "The CSS file ensures the application is visually appealing and responsive:"
    )
    add_bullets(
        doc,
        [
            "CSS Variables: Defines a cohesive color palette (e.g., --bg-color, --text-primary, --accent-color) in the :root for consistency.",
            "Video Wrapper: Uses relative positioning to stack the canvas perfectly over the hidden video stream. The canvas uses a glowing drop-shadow effect to enhance the visibility of the drawn skeleton.",
            "Dashboard Grid: Utilizes CSS Grid (grid-template-columns: repeat(auto-fit, minmax(150px, 1fr))) to automatically adjust the layout of statistics boxes based on screen width.",
            "Typography and Buttons: Implements the 'Outfit' Google Font and styles the calibration button with hover/disabled states, giving it a modern, rounded appearance.",
        ],
    )

    code_heading = doc.add_heading("script.js: Application Logic", level=2)
    code_heading.paragraph_format.page_break_before = True
    doc.add_paragraph(
        "The script.js file encapsulates the core intelligence of the application, managing the camera, AI processing, and state. It is divided into several logical blocks:"
    )
    
    doc.add_heading("Initialization & Variables", level=3)
    doc.add_paragraph(
        "DOM elements (video, canvas, UI counters) are cached in variables. The DEFAULT_BASELINE defines fallback geometric ratios (f1: 0.031, f2: 1.352). SLOUCH_DELTA (-0.12) and LEAN_DELTA (0.55) represent the thresholds for posture deviation. Variables track the current state, session time, and whether the system is calibrated."
    )
    
    doc.add_heading("Geometric Math Helpers", level=3)
    add_bullets(
        doc,
        [
            "distance(p1, p2): Calculates the Euclidean distance between two 2D points.",
            "incenter(a, b, c): Calculates the incenter (center of the incircle) of a triangle formed by three points. This is used to find a stable center point for the face using the eyes, ears, and nose.",
        ],
    )

    doc.add_heading("Validation Functions", level=3)
    add_bullets(
        doc,
        [
            "isVisible(landmark, threshold): Ensures a landmark's confidence score exceeds a required threshold.",
            "isRealHumanUpperBody(landmarks): Acts as a guardrail. It checks for critical joints (nose, eyes, shoulders, ears) with high confidence. It ensures shoulders are below the nose, the shoulder span is realistic (>0.10), and the head-to-shoulder ratio is within plausible human bounds.",
        ],
    )

    doc.add_heading("Posture Prediction", level=3)
    doc.add_paragraph(
        "The predictPosture(f1, f2) function calculates the difference between current measurements and the calibrated baseline. If the slouch difference (f2 - baselineF2) is less than the SLOUCH_DELTA (-0.12), it returns 1 (Slouch). If the lean difference (abs(f1 - baselineF1)) exceeds LEAN_DELTA (0.55), it returns 2 (Lean). Otherwise, it returns 0 (Good posture)."
    )
    
    doc.add_heading("MediaPipe processing (onResults)", level=3)
    doc.add_paragraph(
        "This callback runs on every processed frame:"
    )
    add_bullets(
        doc,
        [
            "Canvas Rendering: Clears the canvas and draws the raw video frame.",
            "Validation Check: If isRealHumanUpperBody fails, it resets the state and prompts the user to show their upper body.",
            "Feature Calculation: Extracts coordinates for the nose, eyes, ears, and shoulders. Calculates the left and right face incenters, and measures the distance from these centers to their respective shoulders. It then computes currentF1 and currentF2.",
            "State Updates: Calls predictPosture and updates the UI status overlay, increments counters (only when transitioning into a bad posture state to avoid continuous counting), and updates colors.",
            "Skeleton Drawing: Uses MediaPipe's drawing_utils to render the POSE_CONNECTIONS skeleton over the user, colored based on their current posture status.",
        ],
    )

    doc.add_heading("Camera and Loop Management", level=3)
    add_bullets(
        doc,
        [
            "initializePose(): Configures the MediaPipe Pose instance and starts a recursive requestAnimationFrame loop (processFrame) that feeds the video element into the Pose model sequentially.",
            "startApplication(): Requests webcam access using navigator.mediaDevices.getUserMedia, starts playback, kicks off the session timer (startSessionTimer), and initializes the AI.",
        ],
    )

    doc.add_heading("Testing Strategy and Limitations", level=1)
    doc.add_paragraph("The application underwent manual testing for camera permissions, state transitions, and responsive layout. Limitations include reliance on third-party CDNs, variable accuracy based on lighting/clothing, and the need for broader clinical validation of the heuristic thresholds.")

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph(
        "The project meets its core goal of private real-time posture feedback in a static web "
        "application. By executing entirely on the client, it completely eliminates privacy risks associated with "
        "cloud-based computer vision. The detailed code breakdown illustrates a robust pipeline from raw pixels "
        "to geometric analysis and actionable UI feedback."
    )

    doc.save(DOCS_DIR / "PROJECT_THESIS_v2.docx")

def main():
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    create_conference_report()
    create_project_thesis()
    print(f"Updated Word reports in {DOCS_DIR}")


if __name__ == "__main__":
    main()
